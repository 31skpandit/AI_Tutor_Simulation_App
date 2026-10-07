"""Split any supported file into pages. Each page is either ready text or an image that needs OCR."""

import io
import re
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from PIL import Image, ImageOps

PART_CHARS = 3000  # DOCX / TXT have no pages: they are split into parts of about this size


@dataclass(frozen=True)
class PageInput:
    page_no: int  # 1-based
    text: str | None = None  # set when text could be read directly
    image_png: bytes | None = None  # set when the page must be OCR'd
    method: str = ""  # text_layer | ocr | docx | text


def count_pages(path: Path, kind: str) -> int:
    if kind == "pdf":
        with pymupdf.open(path) as doc:
            return doc.page_count
    if kind == "image":
        with Image.open(path) as img:
            return getattr(img, "n_frames", 1)
    return len(_split_parts(_read_text(path, kind)))


def iter_pages(path: Path, kind: str, *, min_words: int, max_px: int) -> Iterator[PageInput]:
    if kind == "pdf":
        yield from _pdf_pages(path, min_words, max_px)
    elif kind == "image":
        yield from _image_pages(path, max_px)
    elif kind in {"docx", "text"}:
        for no, part in enumerate(_split_parts(_read_text(path, kind)), start=1):
            yield PageInput(page_no=no, text=part, method=kind)
    else:
        raise ValueError(f"Unsupported kind: {kind}")


def page_image(path: Path, kind: str, page_no: int, max_px: int) -> bytes | None:
    """The page image as sent to OCR (None for DOCX/TXT, which have no page images)."""
    if kind == "pdf":
        with pymupdf.open(path) as doc:
            return _render_pdf_page(doc[page_no - 1], max_px)
    if kind == "image":
        with Image.open(path) as img:
            img.seek(page_no - 1)
            return _to_png(img, max_px)
    return None


def render_page_preview(path: Path, kind: str, page_no: int, max_px: int = 1100) -> bytes | None:
    """PNG preview of a page for the review screen (None for DOCX/TXT)."""
    if kind == "pdf":
        with pymupdf.open(path) as doc:
            return _render_pdf_page(doc[page_no - 1], max_px)
    if kind == "image":
        with Image.open(path) as img:
            img.seek(page_no - 1)
            return _to_png(img, max_px)
    return None


# ---------------------------------------------------------------- PDF


_LIST_START = re.compile(r"^(\(\d{1,2}\)|\d{1,2}\.\s|[*•✱]\s?|\(?[a-d]\)\s)")  # not years: '1969.' is text


def _reflow(block_text: str) -> str:
    """Join the lines of one text block into a paragraph. Justified textbook lines often come out one word per
    line ('Mixed\\nEconomy\\n:'); list items '(1)', '*' still start a new line; 'non-\\nresident' is re-joined."""
    out = ""
    for raw in block_text.splitlines():
        line = " ".join(raw.split())
        if not line:
            continue
        if not out:
            out = line
        elif _LIST_START.match(line):
            out += "\n" + line
        elif out.endswith("-") and line[:1].islower():
            out += line
        else:
            out += " " + line
    return out


def reading_order_text(page: "pymupdf.Page") -> str:
    """Text of a page in the order a person reads it, also for two-column pages.

    Measured on the owner's Class 9 History chapter: PyMuPDF's plain text follows the PDF's internal order, so on
    page 18 the right column came before the left one and a 'Do you know?' box split a sentence. Here blocks are
    sorted: a block across the middle of the page (title, wide box) is read in its place from top to bottom; between
    two such blocks the left column is read fully, then the right column.
    """
    # Works on LINES, not blocks: PyMuPDF can merge two columns that start at the same height into one wide block
    # (found by a test), which would hide the columns.
    width, height = page.rect.width, page.rect.height
    middle = width / 2
    lines = []
    for block in page.get_text("dict")["blocks"]:
        if block.get("type") != 0:
            continue
        for line in block["lines"]:
            text = "".join(span["text"] for span in line["spans"]).strip()
            x0, y0, x1, y1 = line["bbox"]
            if not text or (text.isdigit() and (y1 < height * 0.08 or y0 > height * 0.92)):
                continue  # empty, or a page number in the margin
            lines.append((x0, y0, x1, y1, text))
    lines.sort(key=lambda r: (r[1], r[0]))
    parts: list[str] = []

    def paragraphs(group: list[tuple]) -> None:
        """Lines → rows (same baseline) → paragraphs (split at bigger vertical gaps)."""
        rows: list[list] = []
        for x0, y0, x1, y1, text in sorted(group, key=lambda r: (r[1], r[0])):
            if rows and abs(y0 - rows[-1][1]) < 0.5 * (y1 - y0):
                rows[-1][4].append((x0, text))
                rows[-1][3] = max(rows[-1][3], y1)
            else:
                rows.append([x0, y0, x1, y1, [(x0, text)]])
        block = []
        previous_bottom = None
        previous_x = None
        for row in rows:
            line_height = row[3] - row[1]
            gap = previous_bottom is not None and row[1] - previous_bottom > 0.6 * line_height
            # a paragraph starts with a line indented compared with the line BEFORE it (measured: comparing with the
            # column margin split every line of a 'Do you know?' box, whose lines are all slightly indented)
            indented = previous_x is not None and row[0] - previous_x > 10
            previous_x = row[0]
            if (gap or indented) and block:
                parts.append(_reflow("\n".join(block)))
                block = []
            block.append(" ".join(t for _, t in sorted(row[4])))
            previous_bottom = row[3]
        if block:
            parts.append(_reflow("\n".join(block)))
        group.clear()

    left: list[tuple] = []
    right: list[tuple] = []
    wide: list[tuple] = []
    for line in lines:
        x0, x1 = line[0], line[2]
        if x0 < middle - 15 and x1 > middle + 15:  # spans both columns: finish the columns above it first
            paragraphs(left)
            paragraphs(right)
            wide.append(line)
        else:
            paragraphs(wide)
            (left if (x0 + x1) / 2 < middle else right).append(line)
    paragraphs(wide)
    paragraphs(left)
    paragraphs(right)
    return "\n\n".join(p for p in parts if p.strip())


def _pdf_pages(path: Path, min_words: int, max_px: int) -> Iterator[PageInput]:
    with pymupdf.open(path) as doc:
        for index, page in enumerate(doc):
            text = reading_order_text(page).strip()
            if len(text.split()) >= min_words:
                yield PageInput(page_no=index + 1, text=text, method="text_layer")
            else:
                yield PageInput(page_no=index + 1, image_png=_render_pdf_page(page, max_px), method="ocr")


def _render_pdf_page(page: "pymupdf.Page", max_px: int) -> bytes:
    zoom = max_px / max(page.rect.width, page.rect.height)
    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
    return pixmap.tobytes("png")


# ---------------------------------------------------------------- images


def _image_pages(path: Path, max_px: int) -> Iterator[PageInput]:
    with Image.open(path) as img:
        for frame in range(getattr(img, "n_frames", 1)):
            img.seek(frame)
            yield PageInput(page_no=frame + 1, image_png=_to_png(img, max_px), method="ocr")


def _to_png(img: Image.Image, max_px: int) -> bytes:
    picture = ImageOps.exif_transpose(img.copy())  # phone photos: honour the rotation flag
    picture = picture.convert("RGB")
    picture.thumbnail((max_px, max_px))
    buffer = io.BytesIO()
    picture.save(buffer, format="PNG", optimize=True)
    return buffer.getvalue()


# ---------------------------------------------------------------- DOCX / TXT / MD


def _read_text(path: Path, kind: str) -> str:
    if kind == "text":
        return path.read_text(encoding="utf-8", errors="replace")
    import docx  # python-docx

    document = docx.Document(str(path))
    blocks = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def _split_parts(text: str, size: int = PART_CHARS) -> list[str]:
    paragraphs = [p.strip() for p in text.replace("\r\n", "\n").split("\n\n") if p.strip()]
    parts: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if current and len(current) + len(paragraph) + 2 > size:
            parts.append(current)
            current = paragraph
        else:
            current = f"{current}\n\n{paragraph}" if current else paragraph
    if current:
        parts.append(current)
    return parts
