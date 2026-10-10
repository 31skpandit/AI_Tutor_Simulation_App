"""Ingestion pipeline: file types, de-duplication at every level, review → index, jobs. Offline (fake models)."""

import io
import shutil

import pymupdf
import pytest
from PIL import Image, ImageDraw
from sqlmodel import Session, select

from app.core.config import Settings
from app.db.models import Chunk, Job, Page
from app.ingestion.chunking import chunk_text, text_hash
from app.ingestion.files import guess_metadata, unique_target
from app.ingestion.service import IngestionService
from app.jobs.runner import claim_next, recover_stale, run_until_empty
from tests.conftest import FakeBackend

TEXT_PAGE = (
    "Acids are sour substances. Acids turn blue litmus paper red. Bases are bitter and soapy. "
    "Bases turn red litmus blue. When an acid reacts with a base, salt and water are formed. "
    "This reaction is called neutralisation."
)


def make_text_pdf(path, pages):
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=11)
    doc.save(path)


def png_bytes(label: str) -> bytes:
    image = Image.new("RGB", (400, 200), "white")
    ImageDraw.Draw(image).text((20, 80), label, fill="black")
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    return buffer.getvalue()


def make_scanned_pdf(path, labels):
    doc = pymupdf.open()
    for label in labels:
        page = doc.new_page()
        page.insert_image(page.rect, stream=png_bytes(label))
    doc.save(path)


@pytest.fixture
def env(tmp_path, monkeypatch, make_router):
    import app.ingestion.service as service_module

    source = tmp_path / "source"
    source.mkdir()
    monkeypatch.setattr(service_module, "PROJECT_ROOT", tmp_path)
    # Manual-review behaviour by default; the autopilot has its own tests (tests/test_autopilot.py).
    settings = Settings(
        data_dir=tmp_path / "data",
        source_dir=source,
        require_review=True,
        auto_review_clean=False,
        auto_lesson=False,
    )
    backend = FakeBackend()
    router = make_router(backend)
    service = IngestionService(router.engine, router, settings)
    return service, backend, source


def pages_of(service, document_id):
    return {p.page_no: p for p in service.pages(document_id)}


# ---------------------------------------------------------------- metadata & helpers


def test_guess_metadata_from_filename():
    meta = guess_metadata("Class 9 - Science - 5th Chapter.pdf")
    assert meta == {"class_level": "9", "subject": "Science", "chapter_no": 5, "chapter_title": ""}
    assert guess_metadata("std10_maths_chapter-3.pdf")["subject"] == "Mathematics"
    assert guess_metadata("std10_maths_chapter-3.pdf")["chapter_no"] == 3
    assert guess_metadata("random.pdf")["class_level"] == ""


def test_unique_target_never_overwrites(tmp_path):
    (tmp_path / "a.pdf").write_bytes(b"x")
    assert unique_target(tmp_path, "a.pdf").name == "a (2).pdf"
    assert unique_target(tmp_path, "b.pdf").name == "b.pdf"


def test_chunking_respects_size_and_overlaps():
    text = "\n\n".join(f"Paragraph {i}. " + "word " * 60 for i in range(10))
    chunks = chunk_text(text, max_chars=500, overlap_chars=80)
    assert len(chunks) > 3
    assert all(len(c) <= 500 + 80 + 10 for c in chunks)
    assert chunks[1].split()[0] in chunks[0]  # overlap carried over


def test_text_hash_ignores_whitespace_and_case():
    assert text_hash("Acids  turn\nRED") == text_hash("acids turn red")


# ---------------------------------------------------------------- file types


def test_text_layer_pdf_is_read_without_ocr(env):
    service, backend, source = env
    make_text_pdf(source / "chapter.pdf", [TEXT_PAGE, TEXT_PAGE.replace("Acids", "Metals")])
    result = service.add_file(source / "chapter.pdf")
    assert result.status == "new" and result.document.page_count == 2
    stats = service.extract_document(result.document.id)
    assert stats["text"] == 2 and stats["ocr"] == 0
    assert not backend.calls  # no model used
    assert pages_of(service, result.document.id)[1].method == "text_layer"


def test_scanned_pdf_pages_are_ocrd(env):
    service, backend, source = env
    make_scanned_pdf(source / "scan.pdf", ["page one", "page two"])
    document = service.add_file(source / "scan.pdf").document
    stats = service.extract_document(document.id)
    assert stats["ocr"] == 2
    page = pages_of(service, document.id)[1]
    assert page.method == "ocr" and page.text == backend.ocr_text and page.model_ref == "ollama/vision"
    image_part = backend.calls[0]["messages"][0]["content"][1]
    assert image_part["image_url"]["url"].startswith("data:image/png;base64,")


def test_image_file_is_ocrd(env):
    service, backend, source = env
    (source / "photo.png").write_bytes(png_bytes("textbook photo"))
    document = service.add_file(source / "photo.png").document
    assert document.kind == "image" and document.page_count == 1
    assert service.extract_document(document.id)["ocr"] == 1


def test_docx_and_txt_are_read_directly(env):
    import docx

    service, backend, source = env
    word = docx.Document()
    word.add_paragraph("Chapter 5 Acids, Bases and Salts")
    word.add_paragraph(TEXT_PAGE)
    table = word.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Acid", "HCl"
    word.save(source / "notes.docx")
    (source / "notes.txt").write_text("Plain text notes about salts.", encoding="utf-8")

    report = service.scan_source(enqueue=False)
    assert len(report.new) == 2
    for result in report.new:
        stats = service.extract_document(result.document.id)
        assert stats["text"] == 1
    texts = [p.text for d in service.documents() for p in service.pages(d.id)]
    assert any("Acid | HCl" in t for t in texts) and any("salts" in t for t in texts)
    assert not backend.calls


def test_unsupported_and_lock_files_are_ignored(env):
    service, _, source = env
    (source / "slides.pptx").write_bytes(b"x")
    (source / "~$lock.docx").write_bytes(b"x")
    report = service.scan_source(enqueue=False)
    assert [p.name for p in report.unsupported] == ["slides.pptx"]
    assert not report.new


# ---------------------------------------------------------------- de-duplication


def test_same_file_twice_and_renamed_copy_are_not_reingested(env):
    service, _, source = env
    make_text_pdf(source / "chapter.pdf", [TEXT_PAGE])
    first = service.scan_source(enqueue=False)
    assert len(first.new) == 1
    shutil.copy(source / "chapter.pdf", source / "chapter COPY renamed.pdf")
    second = service.scan_source(enqueue=False)
    assert not second.new and second.known == 1
    assert len(second.duplicates) == 1 and second.duplicates[0].path.name == "chapter COPY renamed.pdf"
    assert len(service.documents()) == 1


def test_upload_of_existing_content_is_skipped_and_not_saved(env):
    service, _, source = env
    make_text_pdf(source / "chapter.pdf", [TEXT_PAGE])
    service.scan_source(enqueue=False)
    data = (source / "chapter.pdf").read_bytes()
    result = service.save_upload("another name.pdf", data)
    assert result.status == "duplicate"
    assert not (source / "another name.pdf").exists()


def test_upload_of_new_content_is_saved_and_registered(env):
    service, _, source = env
    buffer = io.BytesIO()
    make_text_pdf(buffer, [TEXT_PAGE])
    result = service.save_upload("Class 9 - Science - 5th Chapter.pdf", buffer.getvalue())
    assert result.status == "new" and (source / "Class 9 - Science - 5th Chapter.pdf").exists()
    assert result.document.class_level == "9" and result.document.chapter_no == 5
    assert service.save_upload("x.exe", b"1").status == "unsupported"


def test_identical_page_images_are_ocrd_only_once(env):
    service, backend, source = env
    make_scanned_pdf(source / "a.pdf", ["same page", "same page", "different"])
    make_scanned_pdf(source / "b.pdf", ["same page", "another one"])
    for result in service.scan_source(enqueue=False).new:
        service.extract_document(result.document.id)
    # 5 pages, but only 3 distinct images → 3 OCR calls
    assert len(backend.calls) == 3
    methods = [p.method for d in service.documents() for p in service.pages(d.id)]
    assert methods.count("ocr_reused") == 2


def test_extraction_resumes_and_never_repeats_done_pages(env):
    service, backend, source = env
    make_scanned_pdf(source / "scan.pdf", ["one", "two", "three"])
    document = service.add_file(source / "scan.pdf", enqueue=False).document
    service.extract_document(document.id)
    calls = len(backend.calls)
    stats = service.extract_document(document.id)
    assert stats["skipped"] == 3 and len(backend.calls) == calls


def test_failed_ocr_page_is_marked_and_retried(env):
    service, backend, source = env
    make_scanned_pdf(source / "scan.pdf", ["one"])
    document = service.add_file(source / "scan.pdf", enqueue=False).document
    backend.script = [TimeoutError("slow")] * 4  # ocr: call + retry; ocr_retry: call + retry
    stats = service.extract_document(document.id)
    assert stats["failed"] == 1 and service.documents()[0].status == "failed"
    service.retry_page(document.id, 1)  # forgets the page and queues extraction
    stats = service.extract_document(document.id)
    assert stats["ocr"] == 1 and service.documents()[0].status == "review"


def test_duplicate_passages_are_embedded_once(env):
    service, backend, source = env
    make_text_pdf(source / "a.pdf", [TEXT_PAGE])
    make_text_pdf(source / "b.pdf", [TEXT_PAGE + " "])  # different file, same passage
    for result in service.scan_source(enqueue=False).new:
        service.extract_document(result.document.id)
        service.index_document(result.document.id, include_unreviewed=True)
    with Session(service.engine) as s:
        chunks = s.exec(select(Chunk)).all()
    assert len(chunks) == 1
    assert len(backend.embed_calls) == 1


STUCK = "Real first line of the page.\n" + "[Figure: Diagram of the electronic configuration]\n" * 196


def test_repetition_detection():
    from app.ingestion.ocr import collapse_repeats, looks_degenerate

    assert looks_degenerate(STUCK)
    assert not looks_degenerate(TEXT_PAGE * 3)
    assert not looks_degenerate("| a | b |\n| - | - |\n| 1 | 2 |\n")  # small tables are fine
    cleaned = collapse_repeats(STUCK)
    assert cleaned.count("[Figure: Diagram") == 2 and cleaned.startswith("Real first line")


def test_stuck_ocr_is_retried_with_repetition_penalty(env):
    from app.llm.backend import BackendResult

    service, backend, source = env
    make_scanned_pdf(source / "scan.pdf", ["stuck page"])
    document = service.add_file(source / "scan.pdf", enqueue=False).document
    backend.script = [BackendResult(STUCK, 1000, 3000)]  # first attempt loops, retry is clean
    stats = service.extract_document(document.id)
    page = pages_of(service, document.id)[1]
    assert page.text == backend.ocr_text and page.error is None and stats["warnings"] == 0
    assert backend.calls[1]["extra_options"] == {"repeat_penalty": 1.1}


def test_stuck_twice_is_cleaned_and_flagged(env):
    from app.llm.backend import BackendResult

    service, backend, source = env
    make_scanned_pdf(source / "scan.pdf", ["stuck page"])
    document = service.add_file(source / "scan.pdf", enqueue=False).document
    backend.script = [BackendResult(STUCK, 1000, 3000), BackendResult(STUCK, 1000, 3000)]
    stats = service.extract_document(document.id)
    page = pages_of(service, document.id)[1]
    assert page.status == "done" and "repeated" in page.error and stats["warnings"] == 1
    assert page.text.count("[Figure: Diagram") == 2


# ---------------------------------------------------------------- review → index


def test_only_reviewed_pages_are_indexed(env):
    service, _, source = env
    make_text_pdf(source / "chapter.pdf", [TEXT_PAGE, TEXT_PAGE.replace("Acids", "Metals")])
    document = service.add_file(source / "chapter.pdf", enqueue=False).document
    service.extract_document(document.id)
    assert service.documents()[0].status == "review"

    assert service.index_document(document.id)["pages"] == 0  # nothing reviewed yet
    service.update_page(document.id, 1, TEXT_PAGE, reviewed=True)
    assert service.index_document(document.id)["pages"] == 1
    assert service.documents()[0].status == "review"  # page 2 still waiting

    assert service.mark_all_reviewed(document.id) == 1
    service.index_document(document.id)
    assert service.documents()[0].status == "indexed"
    assert service.page_counts(document.id) == {
        "done": 2, "failed": 0, "reviewed": 2, "indexed": 2, "to_check": 0, "auto_reviewed": 0,
    }  # fmt: skip


def test_editing_a_page_replaces_its_passages(env):
    service, _, source = env
    make_text_pdf(source / "chapter.pdf", [TEXT_PAGE])
    document = service.add_file(source / "chapter.pdf", enqueue=False).document
    service.extract_document(document.id)
    service.mark_all_reviewed(document.id)
    service.index_document(document.id)
    service.update_page(document.id, 1, "Corrected text: H2SO4 is sulphuric acid.", reviewed=True)
    assert service.documents()[0].status == "review"
    service.index_document(document.id)
    with Session(service.engine) as s:
        texts = [c.text for c in s.exec(select(Chunk)).all()]
    assert texts == ["Corrected text: H2SO4 is sulphuric acid."]


def test_auto_index_when_review_not_required(env):
    service, _, source = env
    service.settings.require_review = False
    make_text_pdf(source / "chapter.pdf", [TEXT_PAGE])
    service.scan_source()
    run_until_empty(service)
    assert service.documents()[0].status == "indexed"


def test_delete_document_removes_pages_and_passages(env):
    service, _, source = env
    make_text_pdf(source / "chapter.pdf", [TEXT_PAGE])
    document = service.add_file(source / "chapter.pdf", enqueue=False).document
    service.extract_document(document.id)
    service.index_document(document.id, include_unreviewed=True)
    service.delete_document(document.id)
    with Session(service.engine) as s:
        assert not s.exec(select(Chunk)).all() and not s.exec(select(Page)).all()
    assert (source / "chapter.pdf").exists()  # file kept unless asked
    assert service.scan_source(enqueue=False).new  # and can be re-added


# ---------------------------------------------------------------- jobs


def test_jobs_run_in_order_and_are_not_duplicated(env):
    service, _, source = env
    make_text_pdf(source / "chapter.pdf", [TEXT_PAGE])
    document = service.scan_source().new[0].document  # queues an extract job
    service.enqueue("extract", document.id)  # duplicate request → same job
    with Session(service.engine) as s:
        assert len(s.exec(select(Job)).all()) == 1
    assert run_until_empty(service) == 1
    with Session(service.engine) as s:
        job = s.exec(select(Job)).one()
    assert job.status == "done" and job.progress == 1.0 and "1 text" in job.message


def test_job_is_claimed_once(env):
    service, _, _ = env
    service.enqueue("extract", 999)
    assert claim_next(service.engine) is not None
    assert claim_next(service.engine) is None


def test_unknown_job_kinds_are_left_waiting_not_failed(env):
    """An older/newer worker must not grab jobs it cannot run (real incident during Phase 2)."""
    service, _, _ = env
    with Session(service.engine) as s:
        s.add(Job(kind="future_feature", document_id=1))
        s.commit()
    assert claim_next(service.engine) is None
    with Session(service.engine) as s:
        assert s.exec(select(Job)).one().status == "queued"


def test_stale_running_job_is_requeued(env):
    from datetime import timedelta

    from app.db.models import utcnow

    service, _, _ = env
    with Session(service.engine) as s:
        s.add(
            Job(kind="extract", document_id=1, status="running", heartbeat_at=utcnow() - timedelta(hours=1))
        )
        s.commit()
    assert recover_stale(service.engine) == 1


def test_failed_job_records_error(env):
    service, _, _ = env
    service.enqueue("extract", 12345)  # no such document
    run_until_empty(service)
    with Session(service.engine) as s:
        job = s.exec(select(Job)).one()
    assert job.status == "failed" and "12345" in job.error
